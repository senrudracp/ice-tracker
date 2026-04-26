#!/usr/bin/env python3
"""
ICE 3.0 Sprint Tracker — Excel → Supabase Sync Script
Usage: python sync.py [path/to/ICE_Plan_vXX.xlsx]
Defaults to ICE_Plan_v18.xlsx in the same folder.
"""

import sys, json, os
from datetime import datetime, date
from openpyxl import load_workbook
import urllib.request
import urllib.error

# ── Config ────────────────────────────────────────────────────────────────
SUPABASE_URL = "https://toweihaqxumykvdppdqr.supabase.co"
SUPABASE_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InRvd2VpaGFxeHVteWt2ZHBwZHFyIiwicm9sZSI6ImFub24iLCJpYXQiOjE3NzcxOTU1MzYsImV4cCI6MjA5Mjc3MTUzNn0.vQYm9FbYd7DbMWZn3aU4pxpz4DoHaJQljWD-MFKya58"
SHEET_NAME   = "Detailed Sprint Plan"
XLSX_PATH    = sys.argv[1] if len(sys.argv) > 1 else "ICE_Plan_v18.xlsx"

HEADERS = {
    "Content-Type": "application/json",
    "apikey": SUPABASE_KEY,
    "Authorization": f"Bearer {SUPABASE_KEY}",
    "Prefer": "return=minimal"
}

def req(method, path, data=None):
    url = f"{SUPABASE_URL}/rest/v1/{path}"
    body = json.dumps(data).encode() if data else None
    r = urllib.request.Request(url, data=body, headers=HEADERS, method=method)
    try:
        with urllib.request.urlopen(r) as resp:
            raw = resp.read()
            return json.loads(raw) if raw else {}
    except urllib.error.HTTPError as e:
        print(f"  HTTP {e.code}: {e.read().decode()}")
        return None

def fmt_date(v):
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m-%d')
    return None

def get_existing_keys():
    """Get all existing (fr, task_num) keys from Supabase"""
    result = req("GET", "tasks?select=fr,task_num,id")
    if not result: return {}
    return {(r['fr'], r['task_num']): r['id'] for r in result}

def read_excel(path):
    """Read all tasks from Detailed Sprint Plan sheet"""
    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET_NAME]
    tasks = []
    current_sprint = None

    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
        sprint_col = str(row[0] or '').strip()
        fr         = str(row[1] or '').strip()
        task_num   = str(row[2] or '').strip()
        wave       = str(row[3] or '').strip()
        ttype      = str(row[4] or '').strip()
        tdet       = str(row[5] or '').strip()
        est        = row[6]
        start      = row[7]
        end        = row[8]

        # Sprint header row
        if sprint_col.startswith('Sprint') and not fr.startswith('FR-'):
            current_sprint = sprint_col.split('·')[0].strip()
            continue

        if not fr.startswith('FR-') or not task_num:
            continue

        tasks.append({
            'fr':           fr,
            'task_num':     task_num,
            'sprint':       current_sprint or 'Sprint 1',
            'wave':         wave,
            'task_type':    ttype,
            'task_details': tdet[:500],
            'estimate':     float(est) if est else None,
            'planned_start': fmt_date(start),
            'planned_end':   fmt_date(end),
        })

    return tasks

def sync():
    if not os.path.exists(XLSX_PATH):
        print(f"File not found: {XLSX_PATH}")
        sys.exit(1)

    print(f"\n{'='*55}")
    print(f"  ICE 3.0 Sprint Tracker — Sync")
    print(f"  Source : {XLSX_PATH}")
    print(f"  Target : {SUPABASE_URL}")
    print(f"  Time   : {datetime.now().strftime('%d %b %Y %H:%M')}")
    print(f"{'='*55}\n")

    print("Reading Excel...")
    tasks = read_excel(XLSX_PATH)
    print(f"  Found {len(tasks)} tasks in Excel\n")

    print("Fetching existing Supabase tasks...")
    existing = get_existing_keys()
    print(f"  Found {len(existing)} tasks in Supabase\n")

    added = updated = skipped = 0

    for i, t in enumerate(tasks):
        key = (t['fr'], t['task_num'])
        if key in existing:
            # UPDATE only planning fields — never touch status/assignee/actual dates
            task_id = existing[key]
            payload = {
                'sprint':        t['sprint'],
                'wave':          t['wave'],
                'task_type':     t['task_type'],
                'task_details':  t['task_details'],
                'estimate':      t['estimate'],
                'planned_start': t['planned_start'],
                'planned_end':   t['planned_end'],
                'synced_at':     datetime.now().isoformat(),
            }
            result = req("PATCH", f"tasks?fr=eq.{t['fr']}&task_num=eq.{t['task_num']}", payload)
            if result is not None:
                updated += 1
            else:
                print(f"  ⚠  Failed to update {t['fr']} {t['task_num']}")
        else:
            # INSERT new task
            payload = {**t, 'synced_at': datetime.now().isoformat()}
            result = req("POST", "tasks", payload)
            if result is not None:
                added += 1
                print(f"  + {t['fr']} {t['task_num']} — {t['task_details'][:50]}")
            else:
                print(f"  ⚠  Failed to insert {t['fr']} {t['task_num']}")

        # Progress indicator
        if (i+1) % 50 == 0:
            print(f"  ... processed {i+1}/{len(tasks)}")

    # Log the sync
    req("POST", "sync_log", {
        'rows_added': added,
        'rows_updated': updated,
        'rows_skipped': skipped,
        'source': os.path.basename(XLSX_PATH)
    })

    print(f"\n{'='*55}")
    print(f"  Sync complete!")
    print(f"  Added   : {added}")
    print(f"  Updated : {updated}")
    print(f"  Skipped : {skipped}")
    print(f"{'='*55}\n")

if __name__ == "__main__":
    sync()
