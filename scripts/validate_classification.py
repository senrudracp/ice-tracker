#!/usr/bin/env python3
"""
Validates classify_dependency() against the ground-truth audit CSV.
Usage: python validate_classification.py
"""
import csv
import sys
from openpyxl import load_workbook
from classify import classify_dependency

XLSX_PATH = "../ICE_Master_Requirements_v132.xlsx"
SHEET_NAME = "ICE 3.0 — Master Requirements"
AUDIT_CSV = "dependency_classification_audit.csv"

COL_FR = 2
COL_TASK_NUM = 5
COL_DEPENDENT_ON = 10

EXPECTED_MAP = {'READY': 'can_start', 'INTRA_FR': 'dep_on', 'CROSS_FR': 'needs_fr'}


def load_ground_truth(path):
    with open(path, encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
    return {(r['FR'], r['Task']): r for r in rows}


def load_computed(path, sheet_name):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name]
    computed = {}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        fr = str(row[COL_FR] or '').strip()
        task = str(row[COL_TASK_NUM] or '').strip()
        if not fr.startswith('FR-') or not task:
            continue
        raw = row[COL_DEPENDENT_ON]
        readiness, dependent_on = classify_dependency(raw, fr)
        computed[(fr, task)] = (readiness, dependent_on, raw)
    return computed


def main():
    truth = load_ground_truth(AUDIT_CSV)
    computed = load_computed(XLSX_PATH, SHEET_NAME)

    mismatches = []
    missing_in_computed = []

    for key, truth_row in truth.items():
        expected = EXPECTED_MAP[truth_row['Correct_Classification']]
        if key not in computed:
            missing_in_computed.append(key)
            continue
        actual, dependent_on, raw = computed[key]
        if actual != expected:
            mismatches.append((key[0], key[1], expected, actual, raw))

    extra_in_computed = set(computed) - set(truth)

    print(f"Ground truth rows : {len(truth)}")
    print(f"Computed rows     : {len(computed)}")
    print(f"Missing (in truth, not computed): {len(missing_in_computed)}")
    print(f"Extra (in computed, not in truth): {len(extra_in_computed)}")
    print(f"Mismatches        : {len(mismatches)}")
    print()

    if missing_in_computed:
        print("-- Missing from computed --")
        for fr, task in missing_in_computed:
            print(f"  {fr} {task}")

    if mismatches:
        print("-- Mismatches (FR, Task, expected, actual, raw) --")
        for fr, task, expected, actual, raw in mismatches:
            print(f"  {fr} {task}: expected={expected} actual={actual} raw={raw!r}")

    if not mismatches and not missing_in_computed:
        print("Zero mismatches against ground truth.")
        return 0
    return 1


if __name__ == "__main__":
    sys.exit(main())
